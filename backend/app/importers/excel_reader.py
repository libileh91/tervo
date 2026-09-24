"""Read-only Excel/CSV ingestion with physical source coordinates."""
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import csv
import io
import math

import pandas as pd
from openpyxl import load_workbook

from app.importers.format_detector import FormatDetector


def json_value(value):
    if value is None or value is pd.NA or value is pd.NaT or (isinstance(value, float) and not math.isfinite(value)):
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


@dataclass(frozen=True)
class SheetPreview:
    columns: list[str]
    rows: list[dict[str, object]]
    total_rows: int
    source: dict = field(default_factory=dict)
    source_rows: list[int] = field(default_factory=list)

    def to_dict(self):
        return dict(columns=self.columns, rows=self.rows, total_rows=self.total_rows,
                    source=self.source, source_rows=self.source_rows)


class ExcelReader:
    SUPPORTED_SUFFIXES = (".xlsx", ".csv")
    DEFAULT_SHEET = 0

    def sheets(self, path: str | Path) -> list[str]:
        path = Path(path)
        if path.suffix.lower() == '.csv':
            return [path.stem]
        if path.suffix.lower() != '.xlsx':
            raise ValueError("Format non supporté : utiliser .xlsx ou .csv")
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            return book.sheetnames
        finally:
            book.close()

    def read(self, path: str | Path, sheet: int | str = 0, *, header_row: int | None = None,
             encoding: str | None = None, separator: str | None = None,
             source_namespace: str | None = None) -> pd.DataFrame:
        path = Path(path)
        if path.suffix.lower() not in self.SUPPORTED_SUFFIXES:
            raise ValueError("Format non supporté : utiliser .xlsx ou .csv")
        source = dict(file=path.name, namespace=source_namespace)
        if path.suffix.lower() == '.csv':
            raw = path.read_bytes()
            selected_encoding = encoding
            if selected_encoding is None:
                try:
                    raw.decode('utf-8-sig')
                    selected_encoding = 'utf-8-sig'
                except UnicodeDecodeError:
                    selected_encoding = 'latin-1'
            text = raw.decode(selected_encoding)
            if separator is None:
                try:
                    separator = csv.Sniffer().sniff(text[:8192], delimiters=';,\t').delimiter
                except csv.Error as exc:
                    raise ValueError("Séparateur indéterminé : le préciser") from exc
            reader = csv.reader(io.StringIO(text, newline=''), delimiter=separator, strict=True)
            rows, physical = [], []
            start = 1
            for row in reader:
                rows.append(row)
                physical.append(start)
                start = reader.line_num + 1
            source.update(sheet=path.stem, encoding=selected_encoding, separator=separator)
        else:
            book = load_workbook(path, read_only=True, data_only=False)
            try:
                if isinstance(sheet, int):
                    if sheet < 0 or sheet >= len(book.sheetnames):
                        raise ValueError("Feuille introuvable")
                    name = book.sheetnames[sheet]
                else:
                    name = sheet
                if name not in book.sheetnames:
                    raise ValueError("Feuille introuvable")
                # Preserve formulas instead of treating an absent cached result as empty.
                rows = [list(row) for row in book[name].iter_rows(values_only=True)]
                physical = list(range(1, len(rows) + 1))
                source.update(sheet=name)
            finally:
                book.close()
        if not rows:
            raise ValueError("Fichier ou feuille vide")
        if header_row is None:
            aliases = {FormatDetector.header(a) for f, names in FormatDetector.HEADER_SYNONYMS.items()
                       for a in (*names, f.value)}
            scores = [sum(FormatDetector.header(str(v)) in aliases for v in row if v is not None)
                      for row in rows[:30]]
            best = max(scores)
            if best < 2 or scores.count(best) != 1:
                raise ValueError("En-tête indéterminé : préciser header_row (ligne physique)")
            index = scores.index(best)
        else:
            if header_row not in physical:
                raise ValueError("Ligne d'en-tête inexistante")
            index = physical.index(header_row)
        header = list(rows[index])
        while header and header[-1] in (None, ''):
            header.pop()
        columns = [str(v).strip() if v is not None else '' for v in header]
        if not columns or any(not c for c in columns) or len(set(columns)) != len(columns):
            raise ValueError("En-têtes vides ou dupliqués")
        data, numbers = [], []
        for row, number in zip(rows[index + 1:], physical[index + 1:]):
            if all(v is None or v == '' for v in row):
                continue
            if any(v is not None and v != '' for v in row[len(columns):]):
                raise ValueError(f"Ligne {number} : valeurs sans en-tête")
            data.append(row[:len(columns)] + [None] * max(0, len(columns) - len(row)))
            numbers.append(number)
        frame = pd.DataFrame(data, columns=columns, dtype=object)
        source['header_row'] = physical[index]
        frame.attrs.update(source=source, source_rows=numbers)
        return frame

    def preview(self, df: pd.DataFrame, limit: int = 10) -> SheetPreview:
        if limit < 0:
            raise ValueError("Limite négative")
        return SheetPreview(columns=list(df.columns),
            rows=[{k: json_value(v) for k, v in row.items()} for row in df.head(limit).to_dict('records')],
            total_rows=len(df), source=dict(df.attrs.get('source', {})),
            source_rows=df.attrs.get('source_rows', list(range(2, len(df) + 2)))[:limit])
